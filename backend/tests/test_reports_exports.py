"""Reports, exports and court readiness: signed reports (self-signed; an organisation's certificate with an RFC 3161
timestamp from a local `openssl ts` authority; changes after signing are found), the HTML report, report templates,
table exports (paging, formulas defused) and the CASE/UCO (checked with case_validate), STIX 2.1 and MISP exports."""
import asyncio
import csv
import io
import json
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent / 'fixtures'))
from pki import ORG_P12_PASSWORD, TsaServer, make_pki  # noqa: E402

from app.core.config import settings  # noqa: E402
from app.services import report_signing  # noqa: E402

UTC = timezone.utc


def _pdf(text='DFIP test report') -> bytes:
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 720, text)
    c.showPage()
    c.save()
    return buf.getvalue()


@pytest.fixture
def signing(tmp_path, monkeypatch):
    for k, v in {'storage_path': str(tmp_path / 'storage'), 'report_signing_p12': None, 'report_signing_p12_password': None,
                 'report_tsa_url': None, 'report_trust_roots': None, 'report_tsa_trust_roots': None,
                 'data_encryption_key': 'unit-test-data-key-0123456789abcdef'}.items():
        monkeypatch.setattr(settings, k, v)
    return tmp_path


def test_self_signed_report_signature(signing):
    signed, info = report_signing.sign(_pdf(), 'RPT-001', 'Lab A')
    assert info['signed'] and info['self_signed'] and info['level'] == 'PAdES-B-B' and not info['timestamped']
    p12 = signing / 'storage' / 'signing' / 'self-signed.p12'
    assert p12.is_file() and p12.stat().st_mode & 0o777 == 0o600 and not (p12.parent / '.signer.p12').exists()
    v = report_signing.verify(signed)
    s = v['signatures'][0]
    assert v['valid'] and v['error'] is None and len(v['signatures']) == 1
    assert s['intact'] and s['cryptographically_valid'] and s['covers_whole_file'] and s['self_signed']
    assert not s['trusted'] and not s['bottom_line'] and s['timestamp'] is None and 'self-signed' in s['signer']
    # The certificate is made once and used again.
    again, info2 = report_signing.sign(_pdf('second'), 'RPT-002')
    assert info2['signer'] == info['signer'] and report_signing.verify(again)['valid']


def test_changes_after_signing_are_found(signing):
    from pyhanko.pdf_utils import generic
    from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
    signed, _ = report_signing.sign(_pdf(), 'RPT-001')
    # A byte of the signed content changed: the digest no longer matches.
    start = signed.index(b'stream', signed.index(b'/Length')) + 20
    flipped = bytearray(signed)
    flipped[start] ^= 0x01
    v = report_signing.verify(bytes(flipped))
    assert not v['valid'] and not v['signatures'][0]['intact'] and v['error'] == 'A signature is broken or does not match the file'
    # A revision added after the signature (as an editor saving the file would): the signature is intact but no longer covers the file.
    w = IncrementalPdfFileWriter(io.BytesIO(signed))
    w.trailer['/Info'] = w.add_object(generic.DictionaryObject({generic.NameObject('/Title'): generic.TextStringObject('changed')}))
    out = io.BytesIO()
    w.write(out)
    v = report_signing.verify(out.getvalue())
    assert not v['valid'] and v['signatures'][0]['intact'] and not v['signatures'][0]['covers_whole_file']
    assert v['error'] == 'The file was changed after it was signed'
    assert report_signing.verify(_pdf()) == {'valid': False, 'signatures': [], 'error': 'The PDF is not signed'}
    assert report_signing.verify(b'not a pdf')['error'].startswith('Not a readable PDF')


def test_organisation_certificate_with_timestamp(signing, monkeypatch):
    pki = make_pki(signing / 'pki')
    tsa = TsaServer(pki).start()
    try:
        monkeypatch.setattr(settings, 'report_signing_p12', str(pki / 'org.p12'))
        monkeypatch.setattr(settings, 'report_signing_p12_password', ORG_P12_PASSWORD)
        monkeypatch.setattr(settings, 'report_tsa_url', tsa.url)
        signed, info = report_signing.sign(_pdf(), 'RPT-010', 'Lab A')
        assert info['level'] == 'PAdES-B-T' and info['timestamped'] and not info['self_signed'] and info['timestamp_error'] is None
        assert 'DFIP Test Lab report signing' in info['signer'] and tsa.requests >= 1  # pyHanko also asks once to size the signature
        # Without the roots configured the signature checks out but is not trusted.
        s = report_signing.verify(signed)['signatures'][0]
        assert s['intact'] and s['covers_whole_file'] and not s['trusted'] and s['timestamp'] and not s['timestamp']['trusted']
        monkeypatch.setattr(settings, 'report_trust_roots', str(pki / 'ca.crt'))
        monkeypatch.setattr(settings, 'report_tsa_trust_roots', str(pki / 'ca.crt'))
        v = report_signing.verify(signed)
        s = v['signatures'][0]
        assert v['valid'] and s['trusted'] and s['bottom_line'] and not s['self_signed']
        assert s['timestamp']['intact'] and s['timestamp']['valid'] and s['timestamp']['trusted']
        when = datetime.fromisoformat(s['timestamp']['time'])
        assert abs((when - datetime.now(UTC)).total_seconds()) < 300
        # Poppler checks the signature independently of pyHanko.
        f = signing / 'signed.pdf'
        f.write_bytes(signed)
        out = subprocess.run(['pdfsig', str(f)], capture_output=True, text=True).stdout
        assert 'Signature is Valid.' in out and 'DFIP Test Lab report signing' in out
        # The authority unreachable: the report is still signed, without a timestamp, and the reason is kept.
        monkeypatch.setattr(settings, 'report_tsa_url', 'http://127.0.0.1:9/tsa')
        _, info = report_signing.sign(_pdf(), 'RPT-011')
        assert info['level'] == 'PAdES-B-B' and not info['timestamped'] and info['timestamp_error']
    finally:
        tsa.stop()


def test_html_report_is_inert(tmp_path):
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Image, Paragraph, Table
    from PIL import Image as PILImage
    from app.services.report_pdf import _markup, render_html
    PILImage.new('RGB', (8, 6), (20, 120, 200)).save(tmp_path / 'pic.png')
    st = getSampleStyleSheet()
    story = [Paragraph('Digital Forensics Investigation Report', st['Title']),
             Paragraph('Seen &lt;script&gt;alert(1)&lt;/script&gt; in <b>cleanup.ps1</b><br/>next line', st['Normal']),
             Table([['Name', 'Value'], ['<img src=x onerror=alert(1)>', 'a & b']]),
             Image(str(tmp_path / 'pic.png'), 8, 6)]
    html = render_html(story, {'title': 'RPT-001 CASE-1', 'report_id': 'RPT-001', 'classification': 'OFFICIAL <SENSITIVE>',
                               'generated': '2026-10-13 10:00 UTC', 'generated_by': 'Ann "A" <ann>'})
    assert '<h1>Digital Forensics Investigation Report</h1>' in html
    assert '&lt;script&gt;alert(1)&lt;/script&gt; in <b>cleanup.ps1</b><br>next line' in html
    assert '<td>&lt;img src=x onerror=alert(1)&gt;</td><td>a &amp; b</td>' in html
    assert html.count('OFFICIAL &lt;SENSITIVE&gt;') == 2 and 'Ann &quot;A&quot; &lt;ann&gt;' in html
    assert '<img alt="" src="data:image/png;base64,' in html
    assert "default-src 'none'; img-src data:; style-src 'unsafe-inline'" in html
    lower = html.lower()
    assert '<script' not in lower and 'onerror=' not in lower.replace('onerror=alert(1)&gt;', '') and 'http://' not in lower and 'https://' not in lower
    # Markup that is not bold, italic or a line break is dropped, and a stray < never opens a tag.
    assert _markup('<font color="red">x</font> <a href="javascript:x">y</a>') == 'x y'
    assert _markup('<img src=x onerror=alert(1)//') == '&lt;img src=x onerror=alert(1)//'
    assert _markup('<scr<b>ipt>') == 'ipt&gt;' and _markup('<b onclick="x">b</b>') == '<b>b</b>'


def test_report_templates():
    from pydantic import ValidationError
    from app.api.reports import SECTION_IDS, TemplateIn
    from app.services.report_pdf import SECTIONS
    assert len(SECTION_IDS) == len(set(SECTION_IDS)) == len(SECTIONS)
    assert SECTION_IDS[0] == 'evidence' and SECTION_IDS[-1] == 'audit' and {'findings', 'network', 'phones', 'malware'} <= set(SECTION_IDS)
    t = TemplateIn(name='  Brief ', sections=['findings', 'evidence'], cover={'classification': ' OFFICIAL '})
    assert t.sections == ['findings', 'evidence'] and t.cover.classification == 'OFFICIAL' and t.include_screenshots and not t.approved_findings_only
    for bad in (['evidence', 'nope'], ['evidence', 'evidence'], []):
        with pytest.raises(ValidationError):
            TemplateIn(name='Brief', sections=bad)


def test_cells_that_would_run_as_formulas():
    from app.api.exports import flatten, safe_cell
    for v in ('=1+1', '+SUM(A1)', '-2+3', '@cmd', '\tx', '\rx', ' =HYPERLINK("http://x")', '＝1'):
        assert safe_cell(v) == "'" + v
    assert safe_cell('payroll.csv') == 'payroll.csv' and safe_cell('') == '' and safe_cell(-2) == -2 and safe_cell(2.5) == 2.5
    assert safe_cell(True) is True and safe_cell(None) is None and safe_cell(2 ** 60) == str(2 ** 60) and safe_cell('x' * 40000) == 'x' * 32000
    row = flatten({'a': 1, 'nk': 'x', 'pe': {'version_info': {'CompanyName': 'Acme'}, 'imports': {'k': ['A', 'B']}}, 'tags': ['x', None, 3],
                   'files': [{'name': 'a'}], 'deep': {'b': {'c': {'d': {'e': 1}}}}})
    assert row == {'a': 1, 'pe.version_info.CompanyName': 'Acme', 'pe.imports.k': 'A; B', 'tags': 'x; ; 3', 'files': '[{"name": "a"}]',
                   'deep.b.c': '{"d": {"e": 1}}'}


def _table_app():
    from fastapi import FastAPI, Header, HTTPException
    app = FastAPI()
    rows = [{'n': i, 'cell': ['=1+1', '+x', '-y', '@z', 'plain', 'bell\x07'][i % 6], 'meta': {'k': i * 2}} for i in range(23)]

    @app.get('/api/items')
    def items(case_id: str, skip: int = 0, limit: int = 100, authorization: str = Header(None)):
        if authorization != 'Bearer t':
            raise HTTPException(401, 'Authentication required')
        if case_id != 'C1':
            raise HTTPException(404, 'Case not found or access denied')
        return {'data': rows[skip:skip + min(limit, 7)], 'total': len(rows)}  # pages of at most 7 rows

    @app.get('/api/summary')
    def summary():
        return {'counts': {'a': 1}}

    @app.get('/api/plain')
    def plain():
        return [{'x': 1}, {'x': 2, 'y': 'b'}]
    return app


class _Request:
    def __init__(self, app, token='t'):
        self.app, self.headers = app, {'authorization': f'Bearer {token}', 'cookie': 'not forwarded'}


def test_table_export_reads_every_page(tmp_path, monkeypatch):
    from fastapi import HTTPException
    from app.api import exports
    app = _table_app()

    def fetch(path, params, token='t'):
        with open(tmp_path / 'spool.jsonl', 'w', encoding='utf-8') as spool:
            return asyncio.run(exports._fetch_rows(_Request(app, token), path, params, spool))
    columns, n, truncated = fetch('/api/items', [('case_id', 'C1')])
    assert (columns, n, truncated) == (['n', 'cell', 'meta.k'], 23, False)
    exports._write_csv(str(tmp_path / 'spool.jsonl'), columns, str(tmp_path / 'out.csv'))
    data = (tmp_path / 'out.csv').read_bytes()
    assert data.startswith(b'\xef\xbb\xbf')
    got = list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
    assert got[0] == columns and [int(r[0]) for r in got[1:]] == list(range(23))
    assert [r[1] for r in got[1:7]] == ["'=1+1", "'+x", "'-y", "'@z", 'plain', 'bell\x07'] and got[23][2] == '44'
    exports._write_xlsx(str(tmp_path / 'spool.jsonl'), columns, str(tmp_path / 'out.xlsx'), 'items')
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / 'out.xlsx').active
    cells = [c for row in ws.iter_rows() for c in row]
    assert ws.max_row == 24 and not any(c.data_type == 'f' for c in cells)  # no formulas at all
    assert [ws.cell(r, 2).value for r in range(2, 8)] == ["'=1+1", "'+x", "'-y", "'@z", 'plain', 'bell'] and ws.cell(24, 1).value == 22
    # A cap on rows is reported, and lists and errors come through.
    monkeypatch.setattr(exports, 'MAX_ROWS', 10)
    assert fetch('/api/items', [('case_id', 'C1')])[1:] == (10, True)
    assert fetch('/api/plain', [])[:2] == (['x', 'y'], 2)
    for path, params, token, status in (('/api/items', [('case_id', 'C2')], 't', 404), ('/api/items', [('case_id', 'C1')], 'bad', 401),
                                        ('/api/summary', [], 't', 400)):
        with pytest.raises(HTTPException) as e:
            fetch(path, params, token)
        assert e.value.status_code == status
    for bad in ('/api/exports/table', 'https://x.example/api/cases', '/api/../x', '/other', '/api//x'):
        with pytest.raises(HTTPException):
            exports._resource(bad)
    assert exports._resource('/api/forensics/timeline?case_id=C1&skip=5&limit=9&q=a') == ('/api/forensics/timeline', [('case_id', 'C1'), ('q', 'a')])


# A small stand-in for the database: the queries the CASE and indicator exports make.
def _get(doc, key):
    for part in key.split('.'):
        doc = doc.get(part) if isinstance(doc, dict) else None
    return doc


def _match(doc, q):
    for k, v in q.items():
        if k == '$or':
            if not any(_match(doc, x) for x in v):
                return False
            continue
        val = _get(doc, k)
        if isinstance(v, dict) and v and all(op.startswith('$') for op in v):
            for op, arg in v.items():
                ok = {'$in': lambda: val in arg, '$gte': lambda: val is not None and val >= arg, '$gt': lambda: val is not None and val > arg,
                      '$ne': lambda: val != arg, '$exists': lambda: (val is not None) == arg,
                      '$elemMatch': lambda: any(isinstance(x, dict) and _match(x, arg) for x in val or [])}[op]()
                if not ok:
                    return False
        elif not (val == v or (isinstance(val, list) and v in val)):
            return False
    return True


class _Cursor(list):
    def sort(self, key, direction=1):
        for k, d in reversed(key if isinstance(key, list) else [(key, direction)]):
            super().sort(key=lambda x: str(_get(x, k)), reverse=d == -1)
        return self

    def limit(self, n):
        return _Cursor(self[:n])


class _Collection:
    def __init__(self, docs):
        self.docs = docs

    def find(self, q=None, projection=None):
        return _Cursor(d for d in self.docs if _match(d, q or {}))

    def find_one(self, q=None, projection=None, sort=None):
        return next(iter(self.find(q)), None)


class _DB(dict):
    def __missing__(self, name):
        return _Collection([])


def _case_db():
    t = datetime(2026, 9, 1, 14, 0, tzinfo=UTC)
    ev = {'evidence_id': 'EVID-001', 'case_id': 'CASE-1', 'filename': 'laptop.E01', 'size': 1048576, 'format': 'E01',
          'md5': 'a' * 32, 'sha1': 'b' * 40, 'sha256': 'c' * 64}
    return _DB({
        'cases': _Collection([{'case_id': 'CASE-1', 'name': 'Payroll leak', 'description': 'Payroll copied to a USB drive'}]),
        'evidence': _Collection([ev, dict(ev, evidence_id='EVID-002', case_id='CASE-2')]),
        'chain_of_custody': _Collection([
            {'_id': 1, 'case_id': 'CASE-1', 'evidence_id': 'EVID-001', 'action': 'EVIDENCE_REGISTERED', 'timestamp': t, 'user_name': 'Ann'},
            {'_id': 2, 'case_id': 'CASE-1', 'evidence_id': 'EVID-001', 'action': 'EVIDENCE_VERIFIED', 'timestamp': '2026-09-01T15:00:00+00:00',
             'user_name': 'Bo', 'result': 'MATCH'},
            {'_id': 3, 'case_id': 'CASE-1', 'evidence_id': 'EVID-009', 'action': 'EVIDENCE_REGISTERED', 'timestamp': t},
            {'_id': 4, 'case_id': 'CASE-1', 'evidence_id': 'EVID-001', 'action': 'FILE_RECOVERED', 'timestamp': t, 'user_name': 'Ann', 'artifact_id': 'A1',
             'original_path': '/Users/alice/payroll.csv'}]),
        'jobs': _Collection([{'job_id': 'JOB-1', 'type': 'FORENSIC_ANALYSIS', 'status': 'COMPLETED', 'created_at': t, 'started_at': t, 'finished_at': t,
                              'payload': {'evidence_id': 'EVID-001', 'profile': 'standard'}, 'result': {'summary': {'plan': ['fs_index', 'fileproc']}}},
                             {'job_id': 'JOB-2', 'type': 'FORENSIC_ANALYSIS', 'status': 'FAILED', 'payload': {'evidence_id': 'EVID-001'}}]),
        'forensic_artifacts': _Collection([
            {'_id': 10, 'artifact_id': 'A1', 'case_id': 'CASE-1', 'evidence_id': 'EVID-001', 'type': 'RECOVERED_FILE', 'filename': 'payroll.csv',
             'size': 10, 'sha256': 'd' * 64, 'original_path': '/Users/alice/payroll.csv', 'timestamp': t},
            {'_id': 19, 'artifact_id': 'A2', 'case_id': 'CASE-1', 'evidence_id': 'EVID-001', 'type': 'RECOVERED_FILE', 'filename': 'old.docx',
             'size': 20, 'sha256': '6' * 64, 'original_path': '/Users/alice/old.docx', 'timestamp': t},
            {'_id': 11, 'case_id': 'CASE-1', 'evidence_id': 'EVID-001', 'type': 'MALWARE_INDICATOR', 'engine': 'YARA', 'rule': 'Wiper',
             'original_path': '/Users/alice/cleanup.ps1', 'inode': 5, 'volume_id': 'p2048', 'timestamp': t},
            {'_id': 12, 'case_id': 'CASE-1', 'evidence_id': 'EVID-001', 'type': 'HASH_SET_MATCH', 'kind': 'KNOWN_BAD', 'sha256': '1' * 64,
             'original_path': '/Temp/svchost.exe', 'sets': [{'name': 'Lab bad', 'kind': 'KNOWN_BAD'}], 'timestamp': t},
            {'_id': 13, 'case_id': 'CASE-1', 'evidence_id': 'EVID-001', 'type': 'EXECUTABLE', 'path': '/Temp/svchost.exe', 'sha256': '1' * 64,
             'md5': '2' * 32, 'score': 5, 'flags': [{'label': 'Packed (UPX)'}, {'label': 'Name of a system program'}], 'timestamp': t},
            {'_id': 14, 'case_id': 'CASE-1', 'evidence_id': 'EVID-001', 'type': 'EXECUTABLE', 'path': '/Tools/ok.exe', 'sha256': '3' * 64,
             'score': 0, 'vt': {'malicious': 4}, 'timestamp': t},
            {'_id': 15, 'case_id': 'CASE-1', 'evidence_id': 'EVID-001', 'type': 'EXECUTABLE', 'path': '/Tools/fine.exe', 'sha256': '4' * 64, 'score': 1},
            {'_id': 16, 'case_id': 'CASE-1', 'evidence_id': 'EVID-003', 'type': 'HTTP_REQUEST', 'host': 'FileTransfer.example.net:8080',
             'url': "http://FileTransfer.example.net:8080/up?n='1'", 'server': '203.0.113.7', 'client': '10.0.0.5', 'time': t,
             'files': [{'name': 'payroll-export.csv', 'direction': 'request', 'part': True}]},
            {'_id': 17, 'case_id': 'CASE-1', 'evidence_id': 'EVID-003', 'type': 'HTTP_REQUEST', 'host': 'news.example.org', 'files': []},
            {'_id': 18, 'case_id': 'CASE-1', 'type': 'THREAT_INTEL_HASH', 'data': {'indicator': '5' * 64, 'virustotal': {'malicious': 2}}}]),
        'fs_entries': _Collection([{'evidence_id': 'EVID-001', 'inode': 5, 'volume_id': 'p2048', 'path': '/Users/alice/cleanup.ps1',
                                    'sha256': 'e' * 64, 'sha1': 'f' * 40, 'md5': '0' * 32}]),
        'findings': _Collection([{'_id': 20, 'case_id': 'CASE-1', 'status': 'APPROVED', 'title': 'Payroll copied', 'description': 'To USB',
                                  'severity': 'HIGH', 'evidence_ids': ['EVID-001'], 'created_at': t},
                                 {'_id': 21, 'case_id': 'CASE-1', 'status': 'DRAFT', 'title': 'Draft', 'evidence_ids': ['EVID-001'], 'created_at': t}]),
    })


def test_case_uco_export_validates(tmp_path, monkeypatch):
    from app.services import case_uco
    monkeypatch.setattr(case_uco, 'get_db', _case_db)
    doc = case_uco.export('CASE-1', {'name': 'Ann'})
    nodes = doc['@graph']
    types = [n['@type'] for n in nodes]
    assert types[0] == 'case-investigation:Investigation' and types.count('uco-observable:File') == 3  # the evidence and two recovered files
    acts = [n for n in nodes if n['@type'] == 'case-investigation:InvestigativeAction']
    assert len(acts) == 6  # registered, verified, recovered (custody), analysed, recovered (no custody record), exported
    rec = [a for a in acts if a.get('uco-action:result')]
    by_id = {n['@id']: n for n in nodes}
    names = sorted(next(f['uco-observable:fileName'] for f in by_id[a['uco-action:result'][0]['@id']]['uco-core:hasFacet']
                        if 'uco-observable:fileName' in f) for a in rec)
    assert names == ['old.docx', 'payroll.csv'] and sum('uco-action:performer' in a for a in rec) == 1  # Ann recovered payroll.csv
    assert types.count('uco-core:Annotation') == 1 and types.count('case-investigation:ProvenanceRecord') == 1
    assert {n['uco-core:name'] for n in nodes if n['@type'] == 'uco-identity:Identity'} == {'Ann', 'Bo'}
    ev = next(n for n in nodes if n['@type'] == 'uco-observable:File' and any(f.get('uco-observable:fileName') == 'laptop.E01' for f in n['uco-core:hasFacet']))
    hashes = {h['uco-types:hashMethod']: h['uco-types:hashValue']['@value'] for f in ev['uco-core:hasFacet'] for h in f.get('uco-observable:hash', [])}
    assert hashes == {'MD5': 'A' * 32, 'SHA1': 'B' * 40, 'SHA256': 'C' * 64}
    assert len({n['@id'] for n in nodes}) == len(nodes) and case_uco.export('CASE-1', {'name': 'Ann'})['@graph'][1]['@id'] == nodes[1]['@id']
    f = tmp_path / 'case.jsonld'
    f.write_text(json.dumps(doc))
    r = subprocess.run(['case_validate', str(f)], capture_output=True, text=True)
    assert r.returncode == 0 and 'Conforms: True' in r.stdout, r.stdout[-3000:]


def test_indicator_exports(monkeypatch):
    import stix2
    from stix2patterns.validator import run_validator
    from app.services import ioc_export
    monkeypatch.setattr(ioc_export, 'get_db', _case_db)
    items = ioc_export.collect('CASE-1')
    by = {(x['kind'], x['value']): x for x in items}
    assert set(by) == {('file', 'e' * 64), ('file', '1' * 64), ('file', '3' * 64), ('file', '5' * 64),
                       ('url', "http://FileTransfer.example.net:8080/up?n='1'"), ('domain', 'filetransfer.example.net'), ('ipv4', '203.0.113.7')}
    yara = by[('file', 'e' * 64)]
    assert yara['verdict'] == 'malicious-activity' and yara['sha1'] == 'f' * 40 and yara['name'] == 'cleanup.ps1'
    both = by[('file', '1' * 64)]  # in a known-bad set and standing out as a program: one indicator, both reasons
    assert both['verdict'] == 'malicious-activity' and len(both['reasons']) == 2 and both['md5'] == '2' * 32
    assert by[('domain', 'filetransfer.example.net')]['verdict'] == 'anomalous-activity'
    assert [x['verdict'] for x in items] == sorted((x['verdict'] for x in items), key=lambda v: v != 'malicious-activity')
    case = {'case_id': 'CASE-1', 'name': 'Payroll leak'}
    bundle = stix2.parse(ioc_export.stix_bundle(case, items, 'amber', {'name': 'Ann'}), allow_custom=False)
    ind = [o for o in bundle.objects if o.type == 'indicator']
    assert len(ind) == len(items) and all(not run_validator(o.pattern) for o in ind)
    assert all(o.object_marking_refs == [stix2.TLP_AMBER.id] for o in ind)
    url = next(o for o in ind if o.pattern.startswith('[url:'))
    assert url.pattern == "[url:value = 'http://FileTransfer.example.net:8080/up?n=\\'1\\'']"
    report = next(o for o in bundle.objects if o.type == 'report')
    assert set(report.object_refs) == {o.id for o in ind}
    ev = ioc_export.misp_event(case, items, 'clear', {'name': 'Ann'})['Event']
    valid = {'sha256': 'Payload delivery', 'sha1': 'Payload delivery', 'md5': 'Payload delivery', 'filename|sha256': 'Payload delivery',
             'url': 'Network activity', 'domain': 'Network activity', 'ip-dst': 'Network activity'}  # MISP describeTypes pairs
    assert all(valid[a['type']] == a['category'] for a in ev['Attribute'])
    assert {'tlp:white', 'dfip:case="CASE-1"'} == {t['name'] for t in ev['Tag']} and ev['threat_level_id'] == '1'
    assert next(a for a in ev['Attribute'] if a['type'] == 'filename|sha256' and a['value'].startswith('cleanup.ps1|'))['to_ids'] is True
    assert next(a for a in ev['Attribute'] if a['type'] == 'domain')['to_ids'] is False
